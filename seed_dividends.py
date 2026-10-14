"""One-time seed: copy the last known Cash Dividend / Bonus Dividend / EPS per company
from the existing dse_companies.xlsx into dse_dividends.json (never overwrites newer entries).
Usage: python seed_dividends.py [dse_companies.xlsx] [dse_dividends.json]"""
import json, os, sys
from openpyxl import load_workbook

xlsx = sys.argv[1] if len(sys.argv) > 1 else "dse_companies.xlsx"
out  = sys.argv[2] if len(sys.argv) > 2 else "dse_dividends.json"
COLS = {"Cash Dividend": "cash_dividend", "Bonus Dividend": "bonus_stock_dividend", "EPS Basic": "eps_annual"}

state = json.load(open(out, encoding="utf-8")) if os.path.exists(out) else {}
wb = load_workbook(xlsx, read_only=True, data_only=True)
seeded = 0
for ws in wb.worksheets:
    rows = list(ws.iter_rows(values_only=True))
    hdr = next((i for i, r in enumerate(rows[:60]) if r and "Cash Dividend" in [str(c).strip() for c in r if c is not None]), None)
    if hdr is None:
        continue
    names = [str(c).strip() if c is not None else "" for c in rows[hdr]]
    rec = state.setdefault(ws.title, {})
    for col, key in COLS.items():
        if col not in names or rec.get(key):
            continue
        ci = names.index(col)
        vals = [str(r[ci]).strip() for r in rows[hdr + 1:] if len(r) > ci and r[ci] not in (None, "", "--", "-")]
        if vals:
            rec[key] = vals[-1]; seeded += 1
    rec.setdefault("declared", "")
    if not any(rec.get(k) for k in COLS.values()):
        state.pop(ws.title, None)
json.dump(state, open(out, "w", encoding="utf-8"), indent=2, ensure_ascii=False, sort_keys=True)
print(f"Seeded {seeded} values for {len(state)} companies -> {out}")
